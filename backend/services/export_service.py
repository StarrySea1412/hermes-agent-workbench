from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


def generate_project_markdown(project):
    title = project.title or "Untitled report"
    lines = [f"# {title}", ""]
    if project.description:
        lines.extend([project.description.strip(), ""])

    outline = project.outline or _outline_from_text(project.final_content)
    if outline:
        lines.extend(["## 报告大纲", ""])
        lines.extend([f"{index + 1}. {item}" for index, item in enumerate(outline)])
        lines.append("")

    if project.final_content:
        lines.extend(["## 正文", "", project.final_content.strip(), ""])
    else:
        lines.extend(["## 正文", "", "当前项目还没有生成正文。", ""])

    return "\n".join(lines).encode("utf-8")


def generate_project_docx(project):
    doc = Document()
    title = doc.add_heading(project.title or "Untitled report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if project.description:
        doc.add_paragraph(project.description)

    outline = project.outline or _outline_from_text(project.final_content)
    if outline:
        doc.add_heading("报告大纲", level=1)
        for item in outline:
            doc.add_paragraph(str(item), style="List Number")

    doc.add_heading("正文", level=1)
    content = (project.final_content or "当前项目还没有生成正文。").strip()
    for block in content.splitlines():
        stripped = block.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            level = min(stripped.count("#"), 3)
            doc.add_heading(stripped.strip("# "), level=level)
        elif stripped.startswith(("-", "*")):
            doc.add_paragraph(stripped.strip("-* "), style="List Bullet")
        else:
            doc.add_paragraph(stripped)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _outline_from_text(text):
    lines = []
    for line in (text or "").splitlines():
        stripped = line.strip()
        if stripped.startswith("#") or stripped.startswith(("-", "*")):
            lines.append(stripped.strip("#-* "))
    return lines[:10]


def _split_markdown_sections(text):
    sections = []
    current = None
    for line in (text or "").splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("#"):
            if current:
                sections.append(current)
            current = {"title": stripped.strip("# "), "bullets": [], "notes": ""}
            continue
        if current is None:
            current = {"title": "Key points", "bullets": [], "notes": ""}
        current["notes"] += stripped + "\n"
        if stripped.startswith(("-", "*")) or len(current["bullets"]) < 5:
            current["bullets"].append(stripped.strip("-* "))
    if current:
        sections.append(current)
    return sections
